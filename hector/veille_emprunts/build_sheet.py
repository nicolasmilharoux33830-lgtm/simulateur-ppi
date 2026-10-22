import json
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

data = json.load(open(sys.argv[1]))
out = sys.argv[2]

COLS = [
    "Date de détection", "Région", "Département", "Code départemental", "Type de collectivité",
    "Nom officiel de la collectivité", "Population", "Année de référence de la population", "SIREN",
    "Nom de l'interlocuteur", "Fonction de l'interlocuteur", "Téléphone", "E-mail",
    "Nature du besoin d'emprunt", "Montant de l'emprunt, si connu", "Niveau de qualification : A, B ou C",
    "Statut de la consultation", "Date de publication", "Date limite de réponse", "Source",
    "Lien direct vers la source", "Appelé", "Annexes / Informations complémentaires",
]
WIDTHS = [13, 16, 17, 11, 13, 28, 11, 13, 12, 20, 26, 15, 26, 60, 15, 13, 34, 13, 13, 48, 40, 12, 40]
FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name=FONT, size=10)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
MAXROW = 2000


def header(ws, names, widths):
    for i, (n, w) in enumerate(zip(names, widths), start=1):
        c = ws.cell(row=1, column=i, value=n)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"


def d(s):
    return date.fromisoformat(s) if s else None


wb = Workbook()

# Onglet 1 : opportunités
ws = wb.active
ws.title = "Opportunités"
header(ws, COLS, WIDTHS)
for r, o in enumerate(data["opportunites"], start=2):
    vals = [d(o["detection"]), o["region"], o["departement"], o["code_dept"], o["type"], o["nom"],
            o["population"], o["annee_pop"], o["siren"], o["interlocuteur"], o["fonction"], o["tel"],
            o["email"], o["nature"], o["montant"], o["niveau"], o["statut"], d(o["publication"]),
            d(o.get("date_limite")), o["source"], o["lien"], o.get("appele", "Non"), ""]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.alignment = BODY_FONT, WRAP_TOP
    for c in (1, 18, 19):
        ws.cell(row=r, column=c).number_format = "DD/MM/YYYY"
    ws.cell(row=r, column=7).number_format = "# ##0"
    ws.cell(row=r, column=15).number_format = '#,##0.00 "€"'
    ws.cell(row=r, column=21).hyperlink = o["lien"]
    ws.cell(row=r, column=21).font = Font(name=FONT, size=10, color="1155CC", underline="single")
for c in (4, 8, 9, 12):
    for r in range(2, len(data["opportunites"]) + 2):
        ws.cell(row=r, column=c).number_format = "@"
ws.auto_filter.ref = f"A1:W{MAXROW}"

dv_niv = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
dv_app = DataValidation(type="list", formula1='"Non,Oui,À rappeler,Injoignable"', allow_blank=True)
ws.add_data_validation(dv_niv)
ws.add_data_validation(dv_app)
dv_niv.add(f"P2:P{MAXROW}")
dv_app.add(f"V2:V{MAXROW}")

for lvl, color in (("A", "C6EFCE"), ("B", "FFEB9C"), ("C", "E7E6E6")):
    ws.conditional_formatting.add(
        f"P2:P{MAXROW}",
        FormulaRule(formula=[f'$P2="{lvl}"'], fill=PatternFill("solid", fgColor=color)))

# Onglet 2 : journal de veille
wj = wb.create_sheet("Journal de veille")
JCOLS = ["Date de la veille", "Publications couvertes", "Nouvelles opportunités", "Besoins confirmés (A)",
         "Besoins à confirmer (B)", "Signaux commerciaux (C)", "Consultations encore ouvertes (au jour de la veille)",
         "Répartition par région et département", "Sources consultées et volumes analysés",
         "Limites de couverture", "Opportunités déjà suivies mises à jour"]
JW = [13, 22, 13, 12, 12, 12, 16, 34, 70, 70, 24]
header(wj, JCOLS, JW)
for r, j in enumerate(data["journal"], start=2):
    opp = "'Opportunités'"
    vals = [d(j["date"]), j["periode"],
            f"=COUNTIF({opp}!$A:$A,$A{r})",
            f'=COUNTIFS({opp}!$A:$A,$A{r},{opp}!$P:$P,"A")',
            f'=COUNTIFS({opp}!$A:$A,$A{r},{opp}!$P:$P,"B")',
            f'=COUNTIFS({opp}!$A:$A,$A{r},{opp}!$P:$P,"C")',
            j["ouvertes"], j["repartition"], j["sources"], j["limites"], j["suivi"]]
    for c, v in enumerate(vals, start=1):
        cell = wj.cell(row=r, column=c, value=v)
        cell.font, cell.alignment = BODY_FONT, WRAP_TOP
    wj.cell(row=r, column=1).number_format = "DD/MM/YYYY"

# Onglet 3 : actes écartés
we = wb.create_sheet("Actes écartés")
ECOLS = ["Date de détection", "Collectivité", "SIREN", "Département", "Acte analysé",
         "Date de publication", "Motif d'exclusion", "Lien direct vers la source"]
EW = [13, 30, 12, 16, 46, 13, 60, 40]
header(we, ECOLS, EW)
for r, e in enumerate(data["ecartes"], start=2):
    vals = [d(e["detection"]), e["nom"], e["siren"], e["departement"], e["acte"],
            d(e["publication"]), e["motif"], e["lien"]]
    for c, v in enumerate(vals, start=1):
        cell = we.cell(row=r, column=c, value=v)
        cell.font, cell.alignment = BODY_FONT, WRAP_TOP
    for c in (1, 6):
        we.cell(row=r, column=c).number_format = "DD/MM/YYYY"
    we.cell(row=r, column=3).number_format = "@"
    we.cell(row=r, column=8).hyperlink = e["lien"]
    we.cell(row=r, column=8).font = Font(name=FONT, size=10, color="1155CC", underline="single")
we.auto_filter.ref = f"A1:H{MAXROW}"

wb.save(out)
print("ok", out)
